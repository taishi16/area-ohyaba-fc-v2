#!/usr/bin/env python3
"""sheets/*.csv から、Googleドライブへアップロードするだけで使える .xlsx を作る。

Googleスプレッドシートで「ファイル → インポート」を7回やらなくて済むように、
7シートぶんを1つのブックにまとめる。
実行: /home/taishi/sss-agent/.venv/bin/python3 tools/make_xlsx.py
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheet_schema import SHEETS, CHOICES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "sheets")
OUT = os.path.join(SRC, "AREA_OHYABA_FC_サイトデータ.xlsx")

NAVY = "FF00264B"
GOLD = "FFE4C100"
LINE = Side(style="thin", color="FFD5DCE4")

# 列幅（見出し名 → 文字数）
WIDTH = {
    "開催日": 12, "開始時刻": 10, "大会名": 30, "対戦相手": 20, "ホーム／アウェイ": 16,
    "会場": 24, "AREA得点": 10, "相手得点": 10, "状態": 10, "備考": 20, "サンプル": 10,
    "日付": 12, "カテゴリ": 12, "見出し": 34, "概要": 34, "本文": 46,
    "画像ファイル名": 24, "外部リンク": 26, "公開": 10,
    "背番号": 8, "選手名": 14, "ふりがな": 16, "ポジション": 12, "写真ファイル名": 22,
    "生年月日": 12, "身長": 8, "体重": 8, "経歴": 30, "紹介文": 26, "表示": 10, "並び順": 8,
    "氏名": 14, "役職": 14, "指導歴": 26, "指導者資格": 16,
    "会社名・店名": 26, "ロゴファイル名": 22, "ウェブサイトURL": 34, "区分": 22,
    "投稿URL": 46, "メモ": 20,
    "項目": 26, "設定値": 34, "説明": 52,
}

# どのシートのどの列にプルダウンを付けるか
DROPDOWN = {
    "試合": {"状態": "状態", "ホーム／アウェイ": "ホーム／アウェイ", "サンプル": "サンプル"},
    "ニュース": {"公開": "公開"},
    "選手": {"ポジション": "ポジション", "表示": "表示"},
    "スタッフ": {"表示": "表示"},
    "スポンサー": {"表示": "表示", "区分": "区分"},
    "Instagram": {"表示": "表示"},
}

GUIDE = [
    ("AREA OHYABA FC 公式サイト データシート", ""),
    ("", ""),
    ("このブックの7つのシートを直すと、ホームページの中身が変わります。", ""),
    ("HTMLやプログラムを触る必要はありません。", ""),
    ("", ""),
    ("■ やってよいこと", "行を足す／行を消す／中身を書き換える／行を並べ替える"),
    ("■ やってはいけないこと", "1行目（青い見出しの行）を書き換える／シート名を変える"),
    ("", ""),
    ("■ 試合を追加する", "「試合」シートに1行足す。開催日と対戦相手だけでも表示されます。"),
    ("■ 試合結果を入れる", "AREA得点と相手得点を入れるだけ。自動で「結果」に変わります。"),
    ("■ 延期・中止のとき", "「状態」欄から 延期 / 中止 を選ぶ。それ以外は空欄のままでOK。"),
    ("■ 選手を退団扱いにする", "行を消さず、「表示」を「表示しない」にする。"),
    ("■ Instagramを差し替える", "「Instagram」シートの投稿URLを貼り替える。"),
    ("", ""),
    ("■ 反映されないとき", "2〜3分待ってページを再読み込み。急ぐときはメニュー「AREA OHYABA」→「今すぐサイトへ反映する」。"),
    ("■ 入力ミスが不安なとき", "メニュー「AREA OHYABA」→「入力内容をチェックする」"),
    ("", ""),
    ("くわしい手順", "CONTENT_UPDATE_GUIDE.md を見てください。"),
]


def load(name):
    with open(os.path.join(SRC, name + ".csv"), encoding="utf-8-sig", newline="") as f:
        r = csv.reader(f)
        return next(r), [row for row in r if any((c or "").strip() for c in row)]


def build():
    wb = Workbook()

    # ---- はじめに ----
    ws = wb.active
    ws.title = "はじめに"
    ws.sheet_properties.tabColor = GOLD[2:]
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 86
    for i, (a, b) in enumerate(GUIDE, 1):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            ws.cell(i, 1).font = Font(bold=True, size=14, color=NAVY[2:])
        elif a.startswith("■"):
            ws.cell(i, 1).font = Font(bold=True, color=NAVY[2:])

    # ---- データシート ----
    for name, header in SHEETS.items():
        hdr, rows = load(name)
        assert hdr == header, f"{name}: 見出しが定義と違います\n{hdr}\n{header}"
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = NAVY[2:]
        ws.append(header)
        for row in rows:
            ws.append(row)

        # 見出し行の見た目
        fill = PatternFill("solid", fgColor=NAVY)
        for c in range(1, len(header) + 1):
            cell = ws.cell(1, c)
            cell.fill = fill
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = WIDTH.get(header[c - 1], 18)
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        # 本文が入る列は折り返し表示
        for c, h in enumerate(header, 1):
            if h in ("本文", "経歴", "指導歴", "概要", "紹介文", "説明"):
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

        # 罫線
        for r in range(1, ws.max_row + 1):
            for c in range(1, len(header) + 1):
                ws.cell(r, c).border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

        # 日付列の書式
        for c, h in enumerate(header, 1):
            if h in ("開催日", "日付"):
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, c).number_format = "yyyy/mm/dd"

        # プルダウン
        for col_name, choice_key in DROPDOWN.get(name, {}).items():
            if col_name not in header:
                continue
            c = get_column_letter(header.index(col_name) + 1)
            opts = [o for o in CHOICES[choice_key] if o != ""]
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(opts),
                                allow_blank=True, showErrorMessage=False)
            dv.prompt = "%s から選んでください（空欄でも可）" % " / ".join(opts)
            dv.promptTitle = col_name
            ws.add_data_validation(dv)
            dv.add(f"{c}2:{c}500")

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print("作成:", path, os.path.getsize(path), "bytes")
